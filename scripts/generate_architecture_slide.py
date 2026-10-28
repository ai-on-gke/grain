#!/usr/bin/env python3
# Copyright 2026 Google LLC
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""
Generates the Google Slides presentation (.pptx) and SVG architectural diagram
for Grain inside Google (Capsules + Cloud Run Proxy + IAP + Remote Executors).
"""

import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_pptx(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)

    # Background
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 20, 28) # #0f141c
    bg.line.color.rgb = RGBColor(15, 20, 28)

    # Color Palette
    C_BLUE = RGBColor(66, 133, 244)     # Google Blue #4285F4
    C_GREEN = RGBColor(52, 168, 83)     # Google Green / Grain #34A853
    C_YELLOW = RGBColor(251, 188, 5)    # Google Yellow #FBBC05
    C_RED = RGBColor(234, 67, 53)       # Google Red #EA4335
    C_PURPLE = RGBColor(163, 113, 247)  # Purple #a371f7
    C_CARD = RGBColor(24, 32, 44)       # #18202c
    C_CARD_BORDER = RGBColor(43, 56, 76) # #2b384c
    C_TEXT = RGBColor(230, 237, 243)    # #e6edf3
    C_MUTED = RGBColor(139, 148, 158)   # #8b949e
    C_WHITE = RGBColor(255, 255, 255)

    def add_card(x, y, w, h, bg_color=C_CARD, border_color=C_CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # Header section
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p0 = tf.paragraphs[0]
    p0.text = "GRAIN ENTERPRISE ARCHITECTURE | GOOGLE INTERNAL DEPLOYMENT"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_BLUE

    p1 = tf.add_paragraph()
    p1.text = "Grain Inside Google: Cloud Run Proxy & Capsule Topology"
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT

    p2 = tf.add_paragraph()
    p2.text = "Central Grain daemon on dedicated Capsule with outbound-only proxy connection, fronted by Cloud Run + IAP, serving distributed User Capsules."
    p2.font.size = Pt(12)
    p2.font.color.rgb = C_MUTED

    # Section 1: User Capsules (Left)
    col1_x = 0.8
    col_w = 3.6
    add_card(col1_x, 1.65, col_w, 4.3, C_CARD, RGBColor(251, 188, 5))
    
    tb1 = slide.shapes.add_textbox(Inches(col1_x + 0.15), Inches(1.75), Inches(col_w - 0.3), Inches(4.1))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p = tf1.paragraphs[0]
    p.text = "USER CAPSULES (Dev Workstations)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_YELLOW

    p = tf1.add_paragraph()
    p.text = "Developer environments running local agents"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Box 1A: User Capsule 1
    add_card(col1_x + 0.2, 2.35, col_w - 0.4, 1.6, RGBColor(18, 24, 34), C_CARD_BORDER)
    tb_uc1 = slide.shapes.add_textbox(Inches(col1_x + 0.3), Inches(2.45), Inches(col_w - 0.6), Inches(1.4))
    tf_uc1 = tb_uc1.text_frame
    tf_uc1.word_wrap = True
    tf_uc1.margin_left = tf_uc1.margin_top = tf_uc1.margin_right = tf_uc1.margin_bottom = 0
    p = tf_uc1.paragraphs[0]
    p.text = "User Capsule (Engineer A)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT

    p = tf_uc1.add_paragraph()
    p.text = "• grain executor daemon\n• Agent CLIs: agy, claude, codex\n• Ambient credentials & SSO auth\n• Tunneled MCP tool execution"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(201, 209, 217)

    # Box 1B: User Capsule 2
    add_card(col1_x + 0.2, 4.1, col_w - 0.4, 1.6, RGBColor(18, 24, 34), C_CARD_BORDER)
    tb_uc2 = slide.shapes.add_textbox(Inches(col1_x + 0.3), Inches(4.2), Inches(col_w - 0.6), Inches(1.4))
    tf_uc2 = tb_uc2.text_frame
    tf_uc2.word_wrap = True
    tf_uc2.margin_left = tf_uc2.margin_top = tf_uc2.margin_right = tf_uc2.margin_bottom = 0
    p = tf_uc2.paragraphs[0]
    p.text = "User Capsule (Engineer N)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_TEXT

    p = tf_uc2.add_paragraph()
    p.text = "• grain executor (Team B)\n• Custom tools & frameworks\n• Outbound gRPC / WebSocket via IAP\n• Local process execution"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(201, 209, 217)

    # Section 2: Cloud Run Proxy + IAP (Middle)
    col2_x = 4.85
    col2_w = 3.6
    add_card(col2_x, 1.65, col2_w, 4.3, C_CARD, RGBColor(66, 133, 244))

    tb2 = slide.shapes.add_textbox(Inches(col2_x + 0.15), Inches(1.75), Inches(col2_w - 0.3), Inches(4.1))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

    p = tf2.paragraphs[0]
    p.text = "GOOGLE CLOUD RUN PROXY + IAP"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf2.add_paragraph()
    p.text = "Zero-trust ingress & reverse tunnel gateway"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Box 2A: Identity-Aware Proxy (IAP)
    add_card(col2_x + 0.2, 2.35, col2_w - 0.4, 1.6, RGBColor(18, 24, 34), RGBColor(234, 67, 53))
    tb_iap = slide.shapes.add_textbox(Inches(col2_x + 0.3), Inches(2.45), Inches(col2_w - 0.6), Inches(1.4))
    tf_iap = tb_iap.text_frame
    tf_iap.word_wrap = True
    tf_iap.margin_left = tf_iap.margin_top = tf_iap.margin_right = tf_iap.margin_bottom = 0
    p = tf_iap.paragraphs[0]
    p.text = "Identity-Aware Proxy (IAP)"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_RED

    p = tf_iap.add_paragraph()
    p.text = "• Google OAuth / OIDC authentication\n• Zero-Trust access control & IAM\n• mTLS & Device posture checks\n• Strict role-based authorization"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(201, 209, 217)

    # Box 2B: Cloud Run Reverse Proxy
    add_card(col2_x + 0.2, 4.1, col2_w - 0.4, 1.6, RGBColor(18, 24, 34), RGBColor(66, 133, 244))
    tb_cr = slide.shapes.add_textbox(Inches(col2_x + 0.3), Inches(4.2), Inches(col2_w - 0.6), Inches(1.4))
    tf_cr = tb_cr.text_frame
    tf_cr.word_wrap = True
    tf_cr.margin_left = tf_cr.margin_top = tf_cr.margin_right = tf_cr.margin_bottom = 0
    p = tf_cr.paragraphs[0]
    p.text = "Cloud Run Proxy Service"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf_cr.add_paragraph()
    p.text = "• Serverless rendezvous & multiplexer\n• Holds Grain Capsule reverse tunnel\n• Routes executor RPCs to Grain\n• No open public inbound ports to VMs"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(201, 209, 217)

    # Section 3: Grain Capsule (Right)
    col3_x = 8.9
    col3_w = 3.633
    add_card(col3_x, 1.65, col3_w, 4.3, C_CARD, RGBColor(52, 168, 83))

    tb3 = slide.shapes.add_textbox(Inches(col3_x + 0.15), Inches(1.75), Inches(col3_w - 0.3), Inches(4.1))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0

    p = tf3.paragraphs[0]
    p.text = "GRAIN CAPSULE (Central Controller)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    p = tf3.add_paragraph()
    p.text = "Dedicated internal capsule host"
    p.font.size = Pt(10)
    p.font.color.rgb = C_MUTED

    # Box 3A: Grain Daemon
    add_card(col3_x + 0.2, 2.35, col3_w - 0.4, 1.6, RGBColor(18, 24, 34), RGBColor(52, 168, 83))
    tb_gd = slide.shapes.add_textbox(Inches(col3_x + 0.3), Inches(2.45), Inches(col3_w - 0.6), Inches(1.4))
    tf_gd = tb_gd.text_frame
    tf_gd.word_wrap = True
    tf_gd.margin_left = tf_gd.margin_top = tf_gd.margin_right = tf_gd.margin_bottom = 0
    p = tf_gd.paragraphs[0]
    p.text = "Grain Daemon & Orchestrator"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_GREEN

    p = tf_gd.add_paragraph()
    p.text = "• Outbound connection -> Cloud Run\n• Task dispatch & priority queue\n• Autonomous merge queue & CI\n• Multi-team concurrency limiter"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(201, 209, 217)

    # Box 3B: Sandboxes
    add_card(col3_x + 0.2, 4.1, col3_w - 0.4, 1.6, RGBColor(18, 24, 34), RGBColor(163, 113, 247))
    tb_sb = slide.shapes.add_textbox(Inches(col3_x + 0.3), Inches(4.2), Inches(col3_w - 0.6), Inches(1.4))
    tf_sb = tb_sb.text_frame
    tf_sb.word_wrap = True
    tf_sb.margin_left = tf_sb.margin_top = tf_sb.margin_right = tf_sb.margin_bottom = 0
    p = tf_sb.paragraphs[0]
    p.text = "Granule Isolated Sandboxes"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_PURPLE

    p = tf_sb.add_paragraph()
    p.text = "• Hermetic disposable microVMs\n• Executes tunneled MCP tool calls\n• Scoped, ephemeral cloud tokens\n• Zero access to host filesystem"
    p.font.size = Pt(9.5)
    p.font.color.rgb = RGBColor(201, 209, 217)

    # Bottom Protocol & Flow Banner
    add_card(0.8, 6.1, 11.733, 1.0, RGBColor(18, 24, 34), C_CARD_BORDER)
    tb_flow = slide.shapes.add_textbox(Inches(0.95), Inches(6.15), Inches(11.4), Inches(0.9))
    tf_flow = tb_flow.text_frame
    tf_flow.word_wrap = True
    tf_flow.margin_left = tf_flow.margin_top = tf_flow.margin_right = tf_flow.margin_bottom = 0

    p = tf_flow.paragraphs[0]
    p.text = "KEY ARCHITECTURAL FLOWS & SECURITY GUARANTEES"
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = C_BLUE

    p = tf_flow.add_paragraph()
    p.text = "1. Outbound Capsule Dialing: Grain Capsule dials OUT to Cloud Run Proxy (no ingress firewall holes in Google internal network).\n2. IAP Authentication: User Capsules authenticate via Google IAP tokens; Cloud Run validates corporate identity and multiplexes streams.\n3. Split Execution: Agent CLIs run on User Capsules with ambient user auth; tool mutations execute in isolated Granule sandboxes on Grain Capsule."
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(201, 209, 217)

    prs.save(output_path)
    print(f"Saved PPTX slide to {output_path}")

def create_svg(output_path):
    svg = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1920 1080" width="1920" height="1080" style="background:#0f141c; font-family:-apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;">
  <defs>
    <linearGradient id="grad-card" x1="0%" y1="0%" x2="0%" y2="100%">
      <stop offset="0%" stop-color="#1c2534"/>
      <stop offset="100%" stop-color="#141b26"/>
    </linearGradient>
    <linearGradient id="grad-subcard" x1="0%" y1="0%" x2="0%" y2="100%">
      <stop offset="0%" stop-color="#232e40"/>
      <stop offset="100%" stop-color="#18202c"/>
    </linearGradient>
    <linearGradient id="grad-blue" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#4285F4"/>
      <stop offset="100%" stop-color="#1967D2"/>
    </linearGradient>
    <linearGradient id="grad-green" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#34A853"/>
      <stop offset="100%" stop-color="#188038"/>
    </linearGradient>
    <linearGradient id="grad-yellow" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#FBBC05"/>
      <stop offset="100%" stop-color="#F29900"/>
    </linearGradient>
    <linearGradient id="grad-red" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#EA4335"/>
      <stop offset="100%" stop-color="#C5221F"/>
    </linearGradient>
    <linearGradient id="grad-purple" x1="0%" y1="0%" x2="100%" y2="100%">
      <stop offset="0%" stop-color="#a371f7"/>
      <stop offset="100%" stop-color="#8957e5"/>
    </linearGradient>
    <filter id="shadow" x="-5%" y="-5%" width="110%" height="110%">
      <feDropShadow dx="0" dy="8" stdDeviation="12" flood-color="#000000" flood-opacity="0.4"/>
    </filter>
    <marker id="arrow-blue" viewBox="0 0 10 10" refX="6" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse">
      <path d="M 0 1 L 10 5 L 0 9 z" fill="#4285F4"/>
    </marker>
    <marker id="arrow-green" viewBox="0 0 10 10" refX="6" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse">
      <path d="M 0 1 L 10 5 L 0 9 z" fill="#34A853"/>
    </marker>
    <marker id="arrow-yellow" viewBox="0 0 10 10" refX="6" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse">
      <path d="M 0 1 L 10 5 L 0 9 z" fill="#FBBC05"/>
    </marker>
    <marker id="arrow-purple" viewBox="0 0 10 10" refX="6" refY="5" markerWidth="6" markerHeight="6" orient="auto-start-reverse">
      <path d="M 0 1 L 10 5 L 0 9 z" fill="#a371f7"/>
    </marker>
  </defs>

  <!-- Header -->
  <g transform="translate(100, 60)">
    <text x="0" y="0" fill="#4285F4" font-size="16" font-weight="700" letter-spacing="1.5">GRAIN ENTERPRISE ARCHITECTURE | GOOGLE INTERNAL DEPLOYMENT</text>
    <text x="0" y="44" fill="#e6edf3" font-size="34" font-weight="700" letter-spacing="-0.5">Grain Inside Google: Cloud Run Proxy &amp; Capsule Topology</text>
    <text x="0" y="80" fill="#8b949e" font-size="18">Central Grain daemon on dedicated Capsule VM with outbound-only proxy connection, fronted by Cloud Run + IAP, serving distributed User Capsules.</text>
  </g>

  <!-- COLUMN 1: User Capsules (Left) -->
  <g transform="translate(100, 180)" filter="url(#shadow)">
    <!-- Container Card -->
    <rect x="0" y="0" width="510" height="600" rx="16" fill="url(#grad-card)" stroke="#FBBC05" stroke-width="2"/>
    
    <!-- Title Section -->
    <rect x="24" y="24" width="8" height="28" rx="4" fill="#FBBC05"/>
    <text x="44" y="45" fill="#FBBC05" font-size="20" font-weight="700" letter-spacing="0.5">USER CAPSULES</text>
    <text x="44" y="70" fill="#8b949e" font-size="14">Developer environments running remote executors</text>

    <!-- Subcard 1: User Capsule 1 -->
    <g transform="translate(24, 90)">
      <rect x="0" y="0" width="462" height="225" rx="12" fill="url(#grad-subcard)" stroke="#2b384c" stroke-width="1.5"/>
      <circle cx="28" cy="28" r="12" fill="#FBBC05" fill-opacity="0.2"/>
      <text x="28" y="33" fill="#FBBC05" font-size="12" font-weight="700" text-anchor="middle">1</text>
      <text x="50" y="34" fill="#e6edf3" font-size="16" font-weight="600">User Capsule (Engineer Workstation A)</text>
      
      <rect x="24" y="55" width="414" height="42" rx="8" fill="#141b26" stroke="#FBBC05" stroke-opacity="0.4" stroke-width="1"/>
      <text x="38" y="81" fill="#e6edf3" font-size="14" font-weight="600">grain executor <tspan fill="#8b949e" font-weight="400">(Remote Executor Daemon)</tspan></text>

      <text x="24" y="122" fill="#8b949e" font-size="13" font-weight="600">LOCAL AGENT CLIs &amp; CREDENTIALS:</text>
      <g transform="translate(24, 134)">
        <rect x="0" y="0" width="70" height="26" rx="6" fill="#388bfd" fill-opacity="0.15" stroke="#388bfd" stroke-width="1"/>
        <text x="35" y="17" fill="#58a6ff" font-size="12" font-weight="600" text-anchor="middle">agy</text>

        <rect x="80" y="0" width="80" height="26" rx="6" fill="#a371f7" fill-opacity="0.15" stroke="#a371f7" stroke-width="1"/>
        <text x="120" y="17" fill="#bc8cff" font-size="12" font-weight="600" text-anchor="middle">claude</text>

        <rect x="170" y="0" width="80" height="26" rx="6" fill="#2ea043" fill-opacity="0.15" stroke="#2ea043" stroke-width="1"/>
        <text x="210" y="17" fill="#3fb950" font-size="12" font-weight="600" text-anchor="middle">codex</text>
      </g>
      <text x="24" y="185" fill="#c9d1d9" font-size="13">• Ambient Google credentials, gcloud &amp; Corp SSO</text>
      <text x="24" y="205" fill="#c9d1d9" font-size="13">• Outbound connection to Cloud Run via IAP</text>
    </g>

    <!-- Subcard 2: User Capsule N -->
    <g transform="translate(24, 335)">
      <rect x="0" y="0" width="462" height="235" rx="12" fill="url(#grad-subcard)" stroke="#2b384c" stroke-width="1.5"/>
      <circle cx="28" cy="28" r="12" fill="#FBBC05" fill-opacity="0.2"/>
      <text x="28" y="33" fill="#FBBC05" font-size="12" font-weight="700" text-anchor="middle">N</text>
      <text x="50" y="34" fill="#e6edf3" font-size="16" font-weight="600">User Capsule (Engineer Workstation N)</text>

      <rect x="24" y="55" width="414" height="42" rx="8" fill="#141b26" stroke="#FBBC05" stroke-opacity="0.4" stroke-width="1"/>
      <text x="38" y="81" fill="#e6edf3" font-size="14" font-weight="600">grain executor <tspan fill="#8b949e" font-weight="400">(--team infrastructure)</tspan></text>

      <text x="24" y="122" fill="#c9d1d9" font-size="13">• Local environment, toolchains &amp; SDKs</text>
      <text x="24" y="145" fill="#c9d1d9" font-size="13">• Receives task dispatch stream via Cloud Run</text>
      <text x="24" y="168" fill="#c9d1d9" font-size="13">• Drives local CLI subprocesses under ambient user auth</text>
      <text x="24" y="191" fill="#c9d1d9" font-size="13">• MCP tool calls tunnel back to Grain Capsule</text>
      <text x="24" y="214" fill="#8b949e" font-size="12">Security: Zero inbound network access; outbound IAP only</text>
    </g>
  </g>

  <!-- COLUMN 2: Google Cloud Run + IAP (Middle) -->
  <g transform="translate(670, 180)" filter="url(#shadow)">
    <!-- Container Card -->
    <rect x="0" y="0" width="560" height="600" rx="16" fill="url(#grad-card)" stroke="#4285F4" stroke-width="2"/>
    
    <!-- Title Section -->
    <rect x="24" y="24" width="8" height="28" rx="4" fill="#4285F4"/>
    <text x="44" y="45" fill="#4285F4" font-size="20" font-weight="700" letter-spacing="0.5">GOOGLE CLOUD RUN PROXY + IAP</text>
    <text x="44" y="70" fill="#8b949e" font-size="14">Zero-Trust perimeter &amp; reverse tunnel rendezvous</text>

    <!-- Subcard 1: IAP Gate -->
    <g transform="translate(24, 90)">
      <rect x="0" y="0" width="512" height="200" rx="12" fill="url(#grad-subcard)" stroke="#EA4335" stroke-width="1.5"/>
      <rect x="24" y="18" width="10" height="24" rx="3" fill="#EA4335"/>
      <text x="44" y="36" fill="#EA4335" font-size="17" font-weight="700">Identity-Aware Proxy (IAP)</text>
      <text x="44" y="58" fill="#8b949e" font-size="13">Google Corporate Zero-Trust Authentication Layer</text>

      <g transform="translate(24, 75)">
        <text x="0" y="20" fill="#c9d1d9" font-size="13.5">🔒 <tspan font-weight="600" fill="#fff">Google Identity Authentication:</tspan> Validates user / SA OIDC tokens</text>
        <text x="0" y="45" fill="#c9d1d9" font-size="13.5">🛡️ <tspan font-weight="600" fill="#fff">IAM Authorization:</tspan> Grain executor &amp; operator access policies</text>
        <text x="0" y="70" fill="#c9d1d9" font-size="13.5">🌐 <tspan font-weight="600" fill="#fff">mTLS &amp; Context-Aware Access:</tspan> Verified Google device posture</text>
        <text x="0" y="95" fill="#8b949e" font-size="12">Unauthenticated requests dropped at Google edge before reaching proxy</text>
      </g>
    </g>

    <!-- Subcard 2: Cloud Run Proxy Service -->
    <g transform="translate(24, 310)">
      <rect x="0" y="0" width="512" height="260" rx="12" fill="url(#grad-subcard)" stroke="#4285F4" stroke-width="1.5"/>
      <rect x="24" y="18" width="10" height="24" rx="3" fill="#4285F4"/>
      <text x="44" y="36" fill="#4285F4" font-size="17" font-weight="700">Cloud Run Rendezvous Proxy Service</text>
      <text x="44" y="58" fill="#8b949e" font-size="13">Serverless bidirectional gRPC / WebSocket Stream Multiplexer</text>

      <g transform="translate(24, 75)">
        <!-- Internal Flow Box -->
        <rect x="0" y="0" width="464" height="60" rx="8" fill="#141b26" stroke="#4285F4" stroke-opacity="0.4" stroke-width="1"/>
        <text x="16" y="25" fill="#58a6ff" font-size="13" font-weight="600">Reverse Tunnel Session Manager</text>
        <text x="16" y="46" fill="#8b949e" font-size="12">Maintains persistent outbound stream from Grain Capsule</text>

        <text x="0" y="85" fill="#c9d1d9" font-size="13.5">• Terminates inbound connections from User Capsules</text>
        <text x="0" y="110" fill="#c9d1d9" font-size="13.5">• Multiplexes executor RPCs &amp; MCP frames over Grain tunnel</text>
        <text x="0" y="135" fill="#c9d1d9" font-size="13.5">• Zero open inbound ports required on Capsule VMs</text>
        <text x="0" y="160" fill="#34A853" font-size="13" font-weight="600">✓ Highly available, autoscaling Google Cloud Run container</text>
      </g>
    </g>
  </g>

  <!-- COLUMN 3: Grain Capsule (Right) -->
  <g transform="translate(1290, 180)" filter="url(#shadow)">
    <!-- Container Card -->
    <rect x="0" y="0" width="530" height="600" rx="16" fill="url(#grad-card)" stroke="#34A853" stroke-width="2"/>
    
    <!-- Title Section -->
    <rect x="24" y="24" width="8" height="28" rx="4" fill="#34A853"/>
    <text x="44" y="45" fill="#34A853" font-size="20" font-weight="700" letter-spacing="0.5">GRAIN CAPSULE</text>
    <text x="44" y="70" fill="#8b949e" font-size="14">Dedicated VM running central Grain controller</text>

    <!-- Subcard 1: Grain Daemon -->
    <g transform="translate(24, 90)">
      <rect x="0" y="0" width="482" height="240" rx="12" fill="url(#grad-subcard)" stroke="#34A853" stroke-width="1.5"/>
      <circle cx="28" cy="28" r="12" fill="#34A853" fill-opacity="0.2"/>
      <text x="28" y="33" fill="#34A853" font-size="12" font-weight="700" text-anchor="middle">G</text>
      <text x="50" y="34" fill="#e6edf3" font-size="16" font-weight="600">Grain Daemon (Central Orchestrator)</text>
      
      <!-- Outbound Connection Badge -->
      <rect x="24" y="55" width="434" height="42" rx="8" fill="#141b26" stroke="#34A853" stroke-opacity="0.5" stroke-width="1"/>
      <text x="38" y="80" fill="#3fb950" font-size="13.5" font-weight="700">Outbound Dial: <tspan fill="#e6edf3" font-weight="400">Reaches out to Cloud Run Proxy</tspan></text>

      <g transform="translate(24, 115)">
        <text x="0" y="15" fill="#c9d1d9" font-size="13">• Task Backlog, Dependency DAG &amp; Concurrency Scheduler</text>
        <text x="0" y="38" fill="#c9d1d9" font-size="13">• Autonomous Merge Queue &amp; Pull Request Engine</text>
        <text x="0" y="61" fill="#c9d1d9" font-size="13">• Multi-Team Concurrency &amp; Rate Limit Protection</text>
        <text x="0" y="84" fill="#c9d1d9" font-size="13">• Web UI &amp; Director Interactive Planning Dashboard</text>
        <text x="0" y="107" fill="#8b949e" font-size="12">Internal VM networking: Outbound only; no incoming firewall holes</text>
      </g>
    </g>

    <!-- Subcard 2: Granule Sandboxes -->
    <g transform="translate(24, 345)">
      <rect x="0" y="0" width="482" height="225" rx="12" fill="url(#grad-subcard)" stroke="#a371f7" stroke-width="1.5"/>
      <circle cx="28" cy="28" r="12" fill="#a371f7" fill-opacity="0.2"/>
      <text x="28" y="33" fill="#a371f7" font-size="12" font-weight="700" text-anchor="middle">S</text>
      <text x="50" y="34" fill="#e6edf3" font-size="16" font-weight="600">Granule Hermetic Sandboxes (MicroVMs)</text>

      <g transform="translate(24, 60)">
        <rect x="0" y="0" width="434" height="42" rx="8" fill="#141b26" stroke="#a371f7" stroke-opacity="0.4" stroke-width="1"/>
        <text x="16" y="26" fill="#bc8cff" font-size="13.5" font-weight="600">Executes Tunneled MCP Tool Calls</text>

        <text x="0" y="65" fill="#c9d1d9" font-size="13">• run_command, read_file, edit_file isolated in VM</text>
        <text x="0" y="88" fill="#c9d1d9" font-size="13">• Scoped ephemeral tokens; no host filesystem access</text>
        <text x="0" y="111" fill="#c9d1d9" font-size="13">• Native CI workflows run hermetically inside sandbox</text>
        <text x="0" y="134" fill="#3fb950" font-size="12" font-weight="600">✓ Prevents unsafe arbitrary code execution on host</text>
      </g>
    </g>
  </g>

  <!-- Connectors and Flow Arrows -->
  <!-- Arrow 1: User Capsule -> IAP (Yellow / Blue) -->
  <g>
    <path d="M 610 320 L 694 320" fill="none" stroke="#FBBC05" stroke-width="3.5" stroke-dasharray="8,5" marker-end="url(#arrow-yellow)"/>
    <rect x="615" y="292" width="70" height="22" rx="4" fill="#0f141c" stroke="#FBBC05" stroke-width="1"/>
    <text x="650" y="307" fill="#FBBC05" font-size="11" font-weight="700" text-anchor="middle">HTTPS/IAP</text>
  </g>

  <!-- Arrow 2: Grain Capsule -> Cloud Run (Green / Outbound Dial) -->
  <g>
    <path d="M 1290 320 L 1206 320" fill="none" stroke="#34A853" stroke-width="4" marker-end="url(#arrow-green)"/>
    <rect x="1215" y="292" width="130" height="22" rx="4" fill="#0f141c" stroke="#34A853" stroke-width="1"/>
    <text x="1280" y="307" fill="#34A853" font-size="11" font-weight="700" text-anchor="middle">OUTBOUND TUNNEL</text>
  </g>

  <!-- Arrow 3: Bidirectional Multiplexed Tool & Dispatch Stream -->
  <g>
    <path d="M 610 500 L 694 500" fill="none" stroke="#4285F4" stroke-width="3" stroke-dasharray="6,4" marker-end="url(#arrow-blue)"/>
    <path d="M 1206 500 L 1290 500" fill="none" stroke="#a371f7" stroke-width="3" stroke-dasharray="6,4" marker-end="url(#arrow-purple)"/>
    <rect x="880" y="580" width="140" height="24" rx="4" fill="#0f141c" stroke="#4285F4" stroke-width="1"/>
    <text x="950" y="596" fill="#58a6ff" font-size="11" font-weight="600" text-anchor="middle">Bidirectional Multiplex</text>
  </g>

  <!-- Bottom Architecture Flow & Security Summary Banner -->
  <g transform="translate(100, 810)" filter="url(#shadow)">
    <rect x="0" y="0" width="1720" height="210" rx="14" fill="url(#grad-card)" stroke="#2b384c" stroke-width="1.5"/>
    
    <rect x="24" y="20" width="6" height="22" rx="3" fill="#4285F4"/>
    <text x="38" y="37" fill="#4285F4" font-size="16" font-weight="700" letter-spacing="0.5">KEY ARCHITECTURAL HIGHLIGHTS &amp; OPERATIONAL LIFECYCLE</text>

    <!-- 3 Summary Grid Columns -->
    <g transform="translate(24, 55)">
      <!-- Item 1 -->
      <g transform="translate(0, 0)">
        <circle cx="16" cy="16" r="16" fill="#34A853" fill-opacity="0.15"/>
        <text x="16" y="21" fill="#34A853" font-size="14" font-weight="700" text-anchor="middle">1</text>
        <text x="42" y="16" fill="#e6edf3" font-size="15" font-weight="600">Outbound Capsule Tunneling</text>
        <text x="42" y="36" fill="#8b949e" font-size="13">Grain Capsule initiates an outbound connection to Cloud</text>
        <text x="42" y="54" fill="#8b949e" font-size="13">Run Proxy. No internal firewall rules or ingress holes needed.</text>
      </g>

      <!-- Item 2 -->
      <g transform="translate(560, 0)">
        <circle cx="16" cy="16" r="16" fill="#EA4335" fill-opacity="0.15"/>
        <text x="16" y="21" fill="#EA4335" font-size="14" font-weight="700" text-anchor="middle">2</text>
        <text x="42" y="16" fill="#e6edf3" font-size="15" font-weight="600">Zero-Trust IAP Authentication</text>
        <text x="42" y="36" fill="#8b949e" font-size="13">Remote executors on User Capsules authenticate via Google</text>
        <text x="42" y="54" fill="#8b949e" font-size="13">IAP (OIDC tokens). Strict IAM authorization and audit logs.</text>
      </g>

      <!-- Item 3 -->
      <g transform="translate(1120, 0)">
        <circle cx="16" cy="16" r="16" fill="#a371f7" fill-opacity="0.15"/>
        <text x="16" y="21" fill="#a371f7" font-size="14" font-weight="700" text-anchor="middle">3</text>
        <text x="42" y="16" fill="#e6edf3" font-size="15" font-weight="600">Split Execution &amp; Sandboxing</text>
        <text x="42" y="36" fill="#8b949e" font-size="13">Agent CLIs run on User Capsules with ambient credentials,</text>
        <text x="42" y="54" fill="#8b949e" font-size="13">while MCP tool actions execute in Grain's isolated sandboxes.</text>
      </g>

      <!-- Bottom Status Line -->
      <line x1="0" y1="80" x2="1670" y2="80" stroke="#2b384c" stroke-width="1"/>
      <g transform="translate(0, 105)">
        <rect x="0" y="-14" width="80" height="22" rx="4" fill="#238636"/>
        <text x="40" y="1" fill="#fff" font-size="11" font-weight="700" text-anchor="middle">VERIFIED</text>
        <text x="95" y="1" fill="#c9d1d9" font-size="13">Deployment Topology: <tspan fill="#58a6ff">User Capsules (Remote Executors) ⟷ Cloud Run (IAP Proxy) ⟷ Grain Capsule (Daemon &amp; Sandboxes)</tspan></text>
      </g>
    </g>
  </g>
</svg>'''
    with open(output_path, "w") as f:
        f.write(svg)
    print(f"Saved SVG architectural diagram to {output_path}")

if __name__ == "__main__":
    os.makedirs("/home/debian/work/docs/slides/screenshots", exist_ok=True)
    os.makedirs("/home/debian/work/docs/brand", exist_ok=True)
    create_pptx("/home/debian/work/docs/slides/grain-google-architecture.pptx")
    create_svg("/home/debian/work/docs/slides/grain-google-architecture.svg")
    create_svg("/home/debian/work/docs/brand/grain-google-architecture.svg")
